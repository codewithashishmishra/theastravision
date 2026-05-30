"""Parse candidate lists for recruitment campaigns."""

from __future__ import annotations

import csv
import io
import re

from django.core.exceptions import ValidationError
from django.core.validators import validate_email

EMAIL_SPLIT_RE = re.compile(r'[,;\s]+')


def _normalize_row(email: str, first_name: str = '', last_name: str = '') -> dict | None:
    email = (email or '').strip().lower()
    if not email:
        return None
    try:
        validate_email(email)
    except ValidationError:
        return None
    local = email.split('@')[0]
    parts = local.replace('.', ' ').replace('_', ' ').split()
    fn = (first_name or '').strip() or (parts[0].title() if parts else 'Candidate')
    ln = (last_name or '').strip() or (parts[-1].title() if len(parts) > 1 else '')
    return {'email': email, 'first_name': fn, 'last_name': ln}


def parse_emails_text(text: str) -> tuple[list[dict], list[str]]:
    errors: list[str] = []
    rows: list[dict] = []
    seen: set[str] = set()
    for raw in EMAIL_SPLIT_RE.split(text or ''):
        row = _normalize_row(raw)
        if not row:
            continue
        if row['email'] in seen:
            continue
        seen.add(row['email'])
        rows.append(row)
    if not rows and (text or '').strip():
        errors.append('No valid email addresses found.')
    return rows, errors


def parse_csv_file(file_obj) -> tuple[list[dict], list[str]]:
    raw = file_obj.read()
    if isinstance(raw, bytes):
        raw = raw.decode('utf-8-sig', errors='replace')
    reader = csv.DictReader(io.StringIO(raw))
    if not reader.fieldnames:
        return [], ['CSV is empty or missing a header row.']
    normalized = {h.strip().lower().replace(' ', '_'): h for h in reader.fieldnames if h}
    if 'email' not in normalized:
        return [], ['CSV must include an "email" column.']
    rows: list[dict] = []
    errors: list[str] = []
    seen: set[str] = set()
    for line_no, record in enumerate(reader, start=2):
        email = (record.get(normalized['email']) or '').strip()
        fn = record.get(normalized.get('first_name', ''), '') if 'first_name' in normalized else ''
        ln = record.get(normalized.get('last_name', ''), '') if 'last_name' in normalized else ''
        row = _normalize_row(email, str(fn), str(ln))
        if not row:
            if email:
                errors.append(f'Line {line_no}: invalid email ({email}).')
            continue
        if row['email'] in seen:
            errors.append(f'Line {line_no}: duplicate email skipped.')
            continue
        seen.add(row['email'])
        rows.append(row)
    if not rows and not errors:
        errors.append('No valid candidate rows found.')
    return rows, errors


def parse_xlsx_file(file_obj) -> tuple[list[dict], list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        return [], ['openpyxl is required for Excel import.']

    wb = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        return [], ['Excel file is empty.']
    headers = [str(h or '').strip().lower().replace(' ', '_') for h in header]
    if 'email' not in headers:
        return [], ['Excel must include an "email" column.']
    email_idx = headers.index('email')
    fn_idx = headers.index('first_name') if 'first_name' in headers else None
    ln_idx = headers.index('last_name') if 'last_name' in headers else None
    rows: list[dict] = []
    errors: list[str] = []
    seen: set[str] = set()
    for line_no, record in enumerate(rows_iter, start=2):
        if not record:
            continue
        email = str(record[email_idx] or '').strip()
        fn = str(record[fn_idx] or '') if fn_idx is not None and fn_idx < len(record) else ''
        ln = str(record[ln_idx] or '') if ln_idx is not None and ln_idx < len(record) else ''
        row = _normalize_row(email, fn, ln)
        if not row:
            if email:
                errors.append(f'Row {line_no}: invalid email.')
            continue
        if row['email'] in seen:
            continue
        seen.add(row['email'])
        rows.append(row)
    if not rows and not errors:
        errors.append('No valid candidate rows found.')
    return rows, errors


def parse_candidate_file(file_obj) -> tuple[list[dict], list[str]]:
    name = (getattr(file_obj, 'name', '') or '').lower()
    if name.endswith('.xlsx') or name.endswith('.xls'):
        return parse_xlsx_file(file_obj)
    return parse_csv_file(file_obj)
